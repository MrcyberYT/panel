import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io

st.set_page_config(page_title="Student Attendance & Behavior Panel", layout="wide")

# Constants & Quick Mappings
PERIODS = ["P1", "P2", "P3", "P4", "Elective", "P5"]
PAL_MAP = {"P": "Present", "A": "Absent", "L": "Late"}

# Initialize Default Session State
if "behaviors_list" not in st.session_state:
    st.session_state.behaviors_list = [
        "Excellent", "Good", "On Task", 
        "Warning - Talking", "Disruptive", "Needs Improvement"
    ]

if "classes_data" not in st.session_state:
    st.session_state.classes_data = {
        "Class 10A": [
            {"id": "STU-001", "name": "Alice Smith"},
            {"id": "STU-002", "name": "Bob Jones"},
            {"id": "STU-003", "name": "Charlie Brown"}
        ],
        "Class 11B": [
            {"id": "STU-001", "name": "Ethan Hunt"},
            {"id": "STU-002", "name": "Fiona Gallagher"}
        ]
    }

# Ensure structure for student records
if "records" not in st.session_state:
    st.session_state.records = {}
    for cls, students in st.session_state.classes_data.items():
        st.session_state.records[cls] = {}
        for stu in students:
            stu_id = stu["id"]
            st.session_state.records[cls][stu_id] = {
                "name": stu["name"],
                "att": {p: "" for p in PERIODS},
                "behav": {p: [] for p in PERIODS},  # Multiple behaviors as a list
                "notes": ""
            }

# Helper function for quick attendance mapping
def parse_att(val):
    clean = str(val).strip().upper()
    return PAL_MAP.get(clean, val if val in ["Present", "Absent", "Late"] else "")

# --- SIDEBAR CONTROLS ---
st.sidebar.title("🎛️ Controls & Setup")

# Active Class Selector
active_class = st.sidebar.selectbox("Select Class", list(st.session_state.records.keys()))

st.sidebar.markdown("---")
st.sidebar.subheader("➕ Add New Class")
new_cls_name = st.sidebar.text_input("Class Name", key="new_cls")
if st.sidebar.button("Create Class") and new_cls_name:
    if new_cls_name not in st.session_state.records:
        st.session_state.records[new_cls_name] = {}
        st.rerun()

st.sidebar.markdown("---")
st.sidebar.subheader(f"➕ Add Student to {active_class}")
new_stu_name = st.sidebar.text_input("Student Name", key="new_stu")
if st.sidebar.button("Add Student") and new_stu_name:
    count = len(st.session_state.records[active_class]) + 1
    new_id = f"STU-{count:03d}"
    st.session_state.records[active_class][new_id] = {
        "name": new_stu_name,
        "att": {p: "" for p in PERIODS},
        "behav": {p: [] for p in PERIODS},
        "notes": ""
    }
    st.rerun()

st.sidebar.markdown("---")
st.sidebar.subheader("⚙️ Behavior Options")
new_behav_opt = st.sidebar.text_input("New Behavior Tag", key="new_behav")
if st.sidebar.button("Add Tag") and new_behav_opt:
    if new_behav_opt not in st.session_state.behaviors_list:
        st.session_state.behaviors_list.append(new_behav_opt)
        st.rerun()

# --- MAIN PANEL ---
st.title("📋 Attendance & Behavior Panel")
st.caption("Tip: Type **P**, **A**, or **L** in attendance boxes. Select multiple behaviors per period.")

if active_class in st.session_state.records:
    students_dict = st.session_state.records[active_class]
    
    # Selected Period View Mode
    selected_period = st.radio("Select Period to Record:", PERIODS, horizontal=True)

    st.subheader(f"Marking {selected_period} for {active_class}")

    for stu_id, info in list(students_dict.items()):
        with st.expander(f"👤 {info['name']} ({stu_id})", expanded=True):
            col1, col2, col3, col4 = st.columns([2, 4, 3, 1])

            # Quick Attendance Entry (P / A / L)
            with col1:
                curr_att = info["att"].get(selected_period, "")
                att_input = st.text_input(
                    f"Att (P/A/L)", 
                    value=curr_att, 
                    key=f"att_{active_class}_{stu_id}_{selected_period}",
                    max_chars=10
                )
                info["att"][selected_period] = parse_att(att_input)

            # Multiple Behavior Tags Entry
            with col2:
                curr_behavs = info["behav"].get(selected_period, [])
                info["behav"][selected_period] = st.multiselect(
                    "Behaviors (Multiple)",
                    options=st.session_state.behaviors_list,
                    default=curr_behavs,
                    key=f"behav_{active_class}_{stu_id}_{selected_period}"
                )

            # Student Notes
            with col3:
                info["notes"] = st.text_input(
                    "Notes", 
                    value=info["notes"], 
                    key=f"notes_{active_class}_{stu_id}"
                )

            # Delete Student Button
            with col4:
                st.write("")
                st.write("")
                if st.button("🗑️", key=f"del_{active_class}_{stu_id}"):
                    del st.session_state.records[active_class][stu_id]
                    st.rerun()

# --- EXCEL GENERATOR ---
def generate_excel():
    wb = openpyxl.Workbook()
    first = True
    
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    headers = ["Student ID", "Student Name"]
    for p in PERIODS:
        headers.extend([f"{p} Att", f"{p} Behaviors"])
    headers.append("Notes")

    for cls_name, stus in st.session_state.records.items():
        ws = wb.active if first else wb.create_sheet(title=cls_name)
        ws.title = cls_name
        first = False

        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font

        for row_idx, (stu_id, info) in enumerate(stus.items(), start=2):
            row_vals = [stu_id, info["name"]]
            for p in PERIODS:
                row_vals.append(info["att"].get(p, ""))
                # Join multiple behaviors into a comma-separated string for Excel export
                behav_list = info["behav"].get(p, [])
                row_vals.append(", ".join(behav_list) if isinstance(behav_list, list) else str(behav_list))
            row_vals.append(info["notes"])

            ws.append(row_vals)

            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).border = thin_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()

st.sidebar.markdown("---")
excel_file = generate_excel()
st.sidebar.download_button(
    label="📥 Download Excel (.xlsx)",
    data=excel_file,
    file_name="Student_Attendance_and_Behavior_Tracker.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
